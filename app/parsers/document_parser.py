import os
from dataclasses import dataclass, field
from typing import Dict, Tuple
from app.config import settings
from app.utils.logger import get_logger

logger = get_logger(__name__)


class DocumentParsingError(Exception):
    """Raised when parsing fails or input is corrupted."""
    pass


class UnsupportedFileTypeError(Exception):
    """Raised when an unallowed or invalid file type is submitted."""
    pass


class DocumentExceedsSizeLimitError(Exception):
    """Raised when an uploaded document exceeds the configured maximum size."""
    pass


@dataclass
class ParsedDocument:
    text: str
    page_count: int
    format: str
    metadata: Dict[str, str] = field(default_factory=dict)
    is_scanned_or_empty: bool = False


def validate_file_bytes(
    file_bytes: bytes,
    original_filename: str,
) -> Tuple[str, str]:
    """
    Validates file size, extension, and binary signature.
    Returns (normalized_extension, content_type).
    """
    if len(file_bytes) > settings.MAX_UPLOAD_SIZE_BYTES:
        raise DocumentExceedsSizeLimitError(
            f"File exceeds maximum allowed size of {settings.MAX_UPLOAD_SIZE_BYTES} bytes."
        )

    _, ext = os.path.splitext(original_filename.lower())
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Extension '{ext}' is unsupported. Allowed: {settings.ALLOWED_EXTENSIONS}"
        )

    # Magic byte verification
    if ext == ".pdf":
        if not file_bytes.startswith(b"%PDF-"):
            raise UnsupportedFileTypeError("File has .pdf extension but invalid PDF signature.")
        content_type = "application/pdf"
    elif ext == ".docx":
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise UnsupportedFileTypeError("File has .docx extension but invalid ZIP/DOCX signature.")
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif ext == ".txt":
        try:
            file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                file_bytes.decode("latin-1")
            except Exception:
                raise UnsupportedFileTypeError("Text file could not be decoded.")
        content_type = "text/plain"
    else:
        raise UnsupportedFileTypeError(f"Unsupported extension: {ext}")

    return ext, content_type


def parse_document(file_path: str) -> ParsedDocument:
    """
    Parses a supported document file into plain text with metadata.
    Supports TXT, PDF (text-based), and DOCX.
    """
    if not os.path.exists(file_path):
        raise DocumentParsingError(f"File not found: {file_path}")

    _, ext = os.path.splitext(file_path.lower())

    if ext == ".txt":
        return _parse_txt(file_path)
    elif ext == ".pdf":
        return _parse_pdf(file_path)
    elif ext == ".docx":
        return _parse_docx(file_path)
    else:
        raise UnsupportedFileTypeError(f"Parser not configured for extension: {ext}")


def _parse_txt(file_path: str) -> ParsedDocument:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            content = f.read()
    except Exception as e:
        raise DocumentParsingError(f"Failed to read text file: {e}")

    is_empty = len(content.strip()) < 5
    return ParsedDocument(
        text=content,
        page_count=1,
        format="txt",
        is_scanned_or_empty=is_empty,
    )


def _parse_pdf(file_path: str) -> ParsedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise DocumentParsingError("PyMuPDF (fitz) is required for PDF parsing.")

    try:
        doc = fitz.open(file_path)
        pages_text = []
        for i, page in enumerate(doc):
            p_text = page.get_text()
            if p_text:
                pages_text.append(f"--- Page {i + 1} ---\n{p_text}")

        total_text = "\n\n".join(pages_text).strip()
        page_count = len(doc)
        doc.close()
    except Exception as e:
        raise DocumentParsingError(f"Corrupted or invalid PDF document: {e}")

    # Detect if PDF has no extractable text (e.g. scanned image-only PDF)
    is_empty = len(total_text.strip()) < 10
    return ParsedDocument(
        text=total_text,
        page_count=page_count,
        format="pdf",
        is_scanned_or_empty=is_empty,
    )


def _parse_docx(file_path: str) -> ParsedDocument:
    try:
        import docx
    except ImportError:
        raise DocumentParsingError("python-docx is required for DOCX parsing.")

    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        total_text = "\n".join(paragraphs).strip()
    except Exception as e:
        raise DocumentParsingError(f"Corrupted or invalid DOCX document: {e}")

    is_empty = len(total_text.strip()) < 5
    return ParsedDocument(
        text=total_text,
        page_count=1,
        format="docx",
        is_scanned_or_empty=is_empty,
    )
